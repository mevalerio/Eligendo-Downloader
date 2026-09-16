from __future__ import annotations

import csv
import io
import math
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import chain
from pathlib import PurePosixPath
from typing import Any, Iterable

from openpyxl import load_workbook

from .utils import canonical_municipality, clean_text, slug


@dataclass(frozen=True)
class ArchiveRow:
    file_name: str
    row_number: int
    region: str | None
    circoscrizione: str | None
    province: str | None
    municipality: str | None
    municipality_key: str | None
    payload: dict[str, Any]


@dataclass(frozen=True)
class PartyResult:
    election_date: date
    region: str | None
    province: str | None
    municipality: str
    municipality_key: str
    party: str
    party_key: str
    votes: int
    round: int | None
    candidate: str | None
    seats: int | None
    source_file: str
    source_row: int


PARTY_FIELDS = (
    "lista",
    "descrlista",
    "descr_lista",
    "denominazione_lista",
    "descrizione_lista",
    "partito",
    "gruppo",
)
PARTY_VOTE_FIELDS = ("voti_lista", "votilista", "voti")
SEAT_FIELDS = ("seggi_lista", "seggilista", "seggi")


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _normalise_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return clean_text(value)
    if isinstance(value, float) and math.isnan(value):
        return ""
    return value


def _normalise_row(row: dict[Any, Any]) -> dict[str, Any]:
    output: dict[str, Any] = {}
    for key, value in row.items():
        if key is None:
            continue
        normalised_key = slug(str(key))
        if not normalised_key:
            continue
        if isinstance(value, list):
            output[normalised_key] = [_normalise_value(item) for item in value]
        else:
            output[normalised_key] = _normalise_value(value)
    return output


def _first(payload: dict[str, Any], names: tuple[str, ...]) -> str | None:
    for name in names:
        value = payload.get(name)
        if value not in (None, ""):
            return clean_text(str(value))
    return None


def _as_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = clean_text(str(value)).replace(" ", "")
    if "," in text:
        try:
            number = Decimal(text.replace(".", "").replace(",", "."))
        except InvalidOperation:
            return None
        return int(number) if number == number.to_integral_value() else None
    text = text.replace(".", "")
    if text.lstrip("+-").isdigit():
        return int(text)
    return None


def _archive_row(file_name: str, row_number: int, payload: dict[str, Any]) -> ArchiveRow:
    municipality = canonical_municipality(
        _first(
            payload,
            (
                "comune",
                "com",
                "desccomune",
                "denominazione_comune",
                "descrizione_comune",
                "denominazione",
            ),
        )
    )
    return ArchiveRow(
        file_name=file_name,
        row_number=row_number,
        region=_first(
            payload,
            ("regione", "reg", "descregione", "denominazione_regione"),
        ),
        circoscrizione=_first(
            payload,
            (
                "circoscrizione",
                "circoscr",
                "circ_reg",
                "desccirceuropea",
                "denominazione_circoscrizione",
            ),
        ),
        province=_first(
            payload,
            ("provincia", "prov", "descprovincia", "denominazione_provincia"),
        ),
        municipality=municipality,
        municipality_key=slug(municipality) if municipality else None,
        payload=payload,
    )


def _headerless_fields(file_name: str, count: int) -> list[str] | None:
    file_key = slug(file_name)
    if "referendum" in file_key and count == 12:
        return [
            "regione",
            "provincia",
            "comune",
            "num_referendum",
            "quesito",
            "elettori",
            "elettori_maschi",
            "votanti",
            "votanti_maschi",
            "numvotisi",
            "numvotino",
            "schede_bianche",
        ]
    if "comunali" in file_key and count == 16:
        return [
            "regione",
            "provincia",
            "comune",
            "elettoritot",
            "votantitot",
            "skbianche",
            "descrlista",
            "votilista",
            "seggilista",
            "cognome",
            "nome",
            "datanascita",
            "luogonascita",
            "sesso",
            "codtipoeletto",
            "voticand",
        ]
    return None


def _delimited_rows(content: bytes, file_name: str) -> Iterable[ArchiveRow]:
    text = _decode(content)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = None
    reader = (
        csv.reader(io.StringIO(text), dialect=dialect)
        if dialect is not None
        else csv.reader(io.StringIO(text), delimiter=";")
    )
    first_row = next(reader, None)
    if not first_row:
        return
    candidate_headers = [slug(str(value)) for value in first_row]
    known_headers = {
        "regione",
        "provincia",
        "comune",
        "circoscrizione",
        "dataelezione",
        "lista",
        "descrlista",
        "elettori",
        "elettoritot",
        "voti",
        "votilista",
        "voti_lista",
    }
    inferred = _headerless_fields(file_name, len(first_row))
    if sum(header in known_headers for header in candidate_headers) >= 2:
        headers = candidate_headers
        data_rows = reader
        start_row = 2
    elif inferred:
        headers = inferred
        data_rows = chain([first_row], reader)
        start_row = 1
    else:
        headers = candidate_headers
        data_rows = reader
        start_row = 2

    for row_number, values in enumerate(data_rows, start=start_row):
        raw = dict(zip(headers, values, strict=False))
        payload = _normalise_row(raw)
        if payload:
            yield _archive_row(file_name, row_number, payload)


def _xlsx_rows(content: bytes, file_name: str) -> Iterable[ArchiveRow]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        for worksheet in workbook.worksheets:
            values = worksheet.iter_rows(values_only=True)
            headers: list[str] | None = None
            header_row = 0
            for row_number, row_values in enumerate(values, start=1):
                if headers is None:
                    candidate_headers = [
                        slug(str(value)) if value not in (None, "") else ""
                        for value in row_values
                    ]
                    if sum(bool(value) for value in candidate_headers) >= 2:
                        headers = candidate_headers
                        header_row = row_number
                    continue
                mapping = {
                    header: value
                    for header, value in zip(headers, row_values, strict=False)
                    if header
                }
                payload = _normalise_row(mapping)
                if not payload or not any(value not in (None, "") for value in payload.values()):
                    continue
                yield _archive_row(
                    f"{file_name}#{worksheet.title}",
                    row_number,
                    payload,
                )
    finally:
        workbook.close()


def _municipality_level_member(file_name: str) -> bool:
    file_key = slug(file_name)
    excluded_markers = (
        "livsez",
        "liv_sez",
        "preferenze",
        "prefeuropee",
        "sezioni",
        "votanti_varie_ore",
        "votantivarieore",
        "candidatilista",
        "candlista",
        "candidcollegio",
    )
    return not any(marker in file_key for marker in excluded_markers)


def parse_zip_archive(
    content: bytes,
    *,
    max_uncompressed_bytes: int,
    municipality_level_only: bool = False,
) -> list[ArchiveRow]:
    rows: list[ArchiveRow] = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        candidates = []
        total_size = 0
        for info in archive.infolist():
            path = PurePosixPath(info.filename.replace("\\", "/"))
            if path.is_absolute() or ".." in path.parts:
                raise ValueError(f"Unsafe path in ZIP archive: {info.filename}")
            total_size += info.file_size
            if total_size > max_uncompressed_bytes:
                raise ValueError("The ZIP archive exceeds the maximum extracted size.")
            if (
                not info.is_dir()
                and path.suffix.casefold() in {".txt", ".csv", ".xlsx"}
                and (
                    not municipality_level_only
                    or _municipality_level_member(info.filename)
                )
            ):
                candidates.append(info)
        if not candidates:
            raise ValueError("The ZIP archive contains no supported TXT, CSV, or XLSX files.")

        for info in candidates:
            file_content = archive.read(info)
            suffix = PurePosixPath(info.filename).suffix.casefold()
            if suffix == ".xlsx":
                rows.extend(_xlsx_rows(file_content, info.filename))
            else:
                rows.extend(_delimited_rows(file_content, info.filename))
    return rows


def normalise_party_result(row: ArchiveRow, *, election_date: date) -> PartyResult | None:
    if not row.municipality:
        return None
    party = _first(row.payload, PARTY_FIELDS)
    if not party:
        return None
    votes = None
    for field in PARTY_VOTE_FIELDS:
        if field in row.payload:
            votes = _as_int(row.payload[field])
            break
    if votes is None:
        return None
    candidate_parts = [
        _first(row.payload, ("cognome", "cognome_candidato")),
        _first(row.payload, ("nome", "nome_candidato")),
    ]
    candidate = clean_text(" ".join(part for part in candidate_parts if part)) or None
    seats = None
    for field in SEAT_FIELDS:
        if field in row.payload:
            seats = _as_int(row.payload[field])
            break
    return PartyResult(
        election_date=election_date,
        region=row.region,
        province=row.province,
        municipality=row.municipality,
        municipality_key=row.municipality_key or slug(row.municipality),
        party=party,
        party_key=slug(party),
        votes=votes,
        round=_as_int(row.payload.get("turno")),
        candidate=candidate,
        seats=seats,
        source_file=row.file_name,
        source_row=row.row_number,
    )
